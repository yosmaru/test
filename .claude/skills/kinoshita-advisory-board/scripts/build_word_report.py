#!/usr/bin/env python3
"""協議録／事業計画書のWordファイル(.docx)を生成する。

report.json（スキーマは references/word-report.md 参照）を読み、
表紙・クリック可能な目次・協議録本文・完成度スコア・レーダーチャート・
リスク登録簿・記入式の宿題ドリル・前提情報回答シートを、
コンサルティングファーム風の抑制されたデザインで組み立てる。

使い方:
    python build_word_report.py --report report.json --output out.docx

依存: python-docx（なければ `pip install python-docx`）
レーダー画像は radar_chart.py でSVGを作り、cairosvg等でPNGに変換して
report.json の radar_png に渡す。
"""

import argparse
import json
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

# --- デザイントークン（マッキンゼー風の抑制した濃紺＋明るい青アクセント） ---
JP_FONT = "游ゴシック"
NAVY = RGBColor(0x0B, 0x1F, 0x33)     # 見出し・表ヘッダの濃紺
INK = RGBColor(0x22, 0x2A, 0x33)      # 本文
MUTE = RGBColor(0x5A, 0x66, 0x72)     # 補助テキスト
ACCENT = RGBColor(0x1F, 0x6F, 0xE0)   # アクセントの青
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY_HEX = "0B1F33"
ACCENT_HEX = "1F6FE0"
ZEBRA_HEX = "F3F6FA"     # 表の交互背景
BOX_HEX = "EEF3FB"       # 評点ボックス・注記の淡い青
RULE_HEX = "C9D3DD"      # 細い罫線
WRITE_HINT = "（ここに記入してください）"


def _set_jp(run_or_style):
    font = run_or_style.font
    font.name = JP_FONT
    rpr = run_or_style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), JP_FONT)
    rfonts.set(qn("w:ascii"), JP_FONT)
    rfonts.set(qn("w:hAnsi"), JP_FONT)


# OOXML CT_PPr / CT_TcPr / CT_TblPr require their child elements in a strict
# schema order. Blindly .append()-ing a new element after siblings that must
# come later (e.g. w:pBdr after w:spacing, or w:shd after w:tcMar) produces a
# document that fails XSD validation (and can be rejected/repaired by Word).
# These order lists let us always insert a new element at the correct slot.
_PPR_ORDER = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
    "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
    "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN",
    "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing",
    "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment",
    "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange",
]
_TCPR_ORDER = [
    "cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap",
    "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark",
]
_TBLPR_ORDER = [
    "tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
    "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd",
    "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription", "tblPrChange",
]


def _insert_ordered(parent, new_el, order):
    """Insert new_el into parent (a pPr/tcPr/tblPr) at its schema-correct slot."""
    tag = new_el.tag.split("}")[-1]
    try:
        idx = order.index(tag)
    except ValueError:
        parent.append(new_el)
        return
    for child in list(parent):
        ctag = child.tag.split("}")[-1]
        if ctag in order and order.index(ctag) > idx:
            child.addprevious(new_el)
            return
    parent.append(new_el)


def _shade(el, fill_hex):
    """セルや段落に背景色を付ける。"""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill_hex)
    if el.tag.endswith("}tc"):
        _insert_ordered(el.get_or_add_tcPr(), shd, _TCPR_ORDER)
    else:
        _insert_ordered(el, shd, _PPR_ORDER)


def _cell_shade(cell, fill_hex):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill_hex)
    _insert_ordered(cell._tc.get_or_add_tcPr(), shd, _TCPR_ORDER)


def _cell_margins(cell, top=40, bottom=40, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    # CT_TcMar (strict schema) requires this exact child order: top, start, bottom, end.
    for tag, val in (("top", top), ("start", left), ("bottom", bottom), ("end", right)):
        e = OxmlElement(f"w:{tag}")
        e.set(qn("w:w"), str(val))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    _insert_ordered(tcPr, mar, _TCPR_ORDER)


def _cell_text(cell, text, bold=False, size=10, color=None, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color if color else INK
    _set_jp(run)


def _row_height(row, cm):
    tr_pr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(cm * 567)))
    h.set(qn("w:hRule"), "atLeast")
    tr_pr.append(h)


def _table(doc, cols_cm, horizontal_only=True):
    """横罫線のみの抑制した表。マッキンゼー風に縦罫線を消す。"""
    t = doc.add_table(rows=0, cols=len(cols_cm))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tblPr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    # CT_TblBorders (strict schema) requires this exact child order and uses
    # "start"/"end" rather than "left"/"right" for the side borders.
    single_edges = {"top", "bottom", "insideH"}
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        if edge in single_edges:
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), RULE_HEX)
        else:
            e.set(qn("w:val"), "none")
        borders.append(e)
    _insert_ordered(tblPr, borders, _TBLPR_ORDER)
    return t


def _apply_widths(row, cols_cm):
    for cell, w in zip(row.cells, cols_cm):
        cell.width = Cm(w)


def _para(doc, text, size=10.5, bold=False, color=None, space_after=6, space_before=0):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color if color else INK
    _set_jp(run)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.28
    return p


def _bullet(doc, text, size=10.5):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = INK
    _set_jp(run)
    p.paragraph_format.line_spacing = 1.25
    return p


def _bottom_border(paragraph, color_hex, size=6):
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color_hex)
    pbdr.append(bottom)
    _insert_ordered(pPr, pbdr, _PPR_ORDER)


def _bookmark(paragraph, name, bm_id):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bm_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bm_id))
    # w:pPr, if present, must remain the very first child of w:p. Insert the
    # bookmark right after it (or at the start if there is no pPr) rather than
    # blindly at index 0, which would otherwise push it ahead of w:pPr and
    # break the content model.
    p_el = paragraph._p
    pPr = p_el.find(qn("w:pPr"))
    insert_at = 1 if pPr is not None else 0
    p_el.insert(insert_at, start)
    p_el.append(end)


def _section_heading(doc, text, num=None, bm=None, bm_id=0):
    """レポートの主要セクション見出し（番号＋濃紺＋下線アクセント）。"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    if num:
        rn = p.add_run(f"{num}  ")
        rn.bold = True
        rn.font.size = Pt(15)
        rn.font.color.rgb = ACCENT
        _set_jp(rn)
    rt = p.add_run(text)
    rt.bold = True
    rt.font.size = Pt(15)
    rt.font.color.rgb = NAVY
    _set_jp(rt)
    _bottom_border(p, RULE_HEX, size=6)
    if bm:
        _bookmark(p, bm, bm_id)
    # アウトラインレベル1（Word標準の見出し扱い）
    pPr = p._p.get_or_add_pPr()
    ol = OxmlElement("w:outlineLvl")
    ol.set(qn("w:val"), "0")
    pPr.append(ol)
    return p


def _sub_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(11.5)
    run.font.color.rgb = NAVY
    _set_jp(run)
    pPr = p._p.get_or_add_pPr()
    ol = OxmlElement("w:outlineLvl")
    ol.set(qn("w:val"), "1")
    pPr.append(ol)
    return p


def _toc_link(doc, num, text, anchor):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    rn = p.add_run(f"{num}   ")
    rn.bold = True
    rn.font.size = Pt(10.5)
    rn.font.color.rgb = ACCENT
    _set_jp(rn)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:eastAsia"), JP_FONT)
    rfonts.set(qn("w:ascii"), JP_FONT)
    rfonts.set(qn("w:hAnsi"), JP_FONT)
    rPr.append(rfonts)
    col = OxmlElement("w:color")
    col.set(qn("w:val"), NAVY_HEX)
    rPr.append(col)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "21")
    rPr.append(sz)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    hyperlink.append(r)
    p._p.append(hyperlink)


def _page_number_footer(doc):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), kind)
            run._r.append(fld)
        else:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = txt
            run._r.append(instr)
    run.font.size = Pt(9)
    run.font.color.rgb = MUTE
    _set_jp(run)


def _l1_title(sec, has):
    """このsectionが目次に載る主要見出しなら、その表示名を返す。"""
    m = sec.get("type")
    labels = {
        "scorecard": ("完成度スコア", "score_table"),
        "risk": ("リスク登録簿", "risk_table"),
        "homework": ("宿題ドリル", "homework"),
        "premise": ("前提情報回答シート", "premise_questions"),
    }
    if m in labels:
        label, key = labels[m]
        return label if has.get(key) else None
    if m:
        return None
    if sec.get("level", 1) == 1 and sec.get("heading"):
        return sec["heading"]
    return None


def build(report, output):
    doc = Document()
    _set_jp(doc.styles["Normal"])
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Normal"].font.color.rgb = INK
    for s in doc.sections:
        s.page_width, s.page_height = Mm(210), Mm(297)
        s.left_margin = s.right_margin = Mm(22)
        s.top_margin, s.bottom_margin = Mm(22), Mm(20)

    is_plan = report.get("doc_type") == "plan"
    ov = report.get("overall") or {}
    has = {
        "score_table": bool(report.get("score_table")),
        "risk_table": bool(report.get("risk_table")),
        "homework": bool(report.get("homework")),
        "premise_questions": bool(report.get("premise_questions")),
    }
    sections = report.get("sections", [])

    # 目次に載る主要見出しを、bookmark名つきで先に確定する
    toc_entries = []  # (num_str, title, bookmark)
    n = 0
    for sec in sections:
        title = _l1_title(sec, has)
        if title:
            n += 1
            toc_entries.append((f"{n:02d}", title, f"bm{n}"))
    bm_iter = iter(toc_entries)

    # ---- 表紙 ----
    top = doc.add_paragraph()
    top.paragraph_format.space_after = Pt(0)
    _bottom_border(top, ACCENT_HEX, size=18)
    doc.add_paragraph().paragraph_format.space_after = Pt(60)

    label = doc.add_paragraph()
    r = label.add_run("ADVISORY REPORT" if not is_plan else "BUSINESS PLAN")
    r.font.size = Pt(10)
    r.font.color.rgb = ACCENT
    r.bold = True
    _set_jp(r)
    label.paragraph_format.space_after = Pt(4)

    title = report.get("title", "伴走型稼ぐまちづくりアドバイザリーボード 協議録")
    tp = doc.add_paragraph()
    tr = tp.add_run(title)
    tr.bold = True
    tr.font.size = Pt(26)
    tr.font.color.rgb = NAVY
    _set_jp(tr)
    tp.paragraph_format.space_after = Pt(6)

    sub = doc.add_paragraph()
    sr = sub.add_run(report.get("case_name", ""))
    sr.font.size = Pt(13)
    sr.font.color.rgb = MUTE
    _set_jp(sr)
    sub.paragraph_format.space_after = Pt(40)

    # メタ情報（罫線の下）
    meta_line = doc.add_paragraph()
    _bottom_border(meta_line, RULE_HEX, size=4)
    meta_line.paragraph_format.space_after = Pt(8)
    metas = []
    if not is_plan:
        metas.append(("協議回", f"第{report.get('session', 1)}回"))
    metas.append(("発行日", report.get("date", "")))
    if ov:
        delta = ov.get("delta")
        dtxt = f"（前回比 {'+' if isinstance(delta,(int,float)) and delta >= 0 else ''}{delta}）" if isinstance(delta, (int, float)) else ""
        metas.append(("完成度", f"総合 {ov.get('score','')}点／{ov.get('stage','')}{dtxt}"))
    for k, v in metas:
        mp = doc.add_paragraph()
        mp.paragraph_format.space_after = Pt(2)
        rk = mp.add_run(f"{k}　")
        rk.bold = True
        rk.font.size = Pt(10)
        rk.font.color.rgb = NAVY
        _set_jp(rk)
        rv = mp.add_run(str(v))
        rv.font.size = Pt(10)
        rv.font.color.rgb = INK
        _set_jp(rv)

    doc.add_page_break()

    # ---- 目次 ----
    toch = doc.add_paragraph()
    tr = toch.add_run("目次")
    tr.bold = True
    tr.font.size = Pt(18)
    tr.font.color.rgb = NAVY
    _set_jp(tr)
    ten = toch.add_run("  CONTENTS")
    ten.font.size = Pt(11)
    ten.font.color.rgb = ACCENT
    ten.bold = True
    _set_jp(ten)
    _bottom_border(toch, RULE_HEX, size=6)
    toch.paragraph_format.space_after = Pt(10)
    for num, title, anchor in toc_entries:
        _toc_link(doc, num, title, anchor)
    _para(doc, "（各項目をクリックすると本文へ移動します）", size=8.5, color=MUTE, space_before=8)
    doc.add_page_break()

    # ---- ブロック描画関数（マーカーで任意位置に配置） ----
    def _next_bm():
        try:
            num, _title, anchor = next(bm_iter)
            return num, anchor
        except StopIteration:
            return None, None

    def render_scorecard():
        if not has["score_table"]:
            return
        num, anchor = _next_bm()
        _section_heading(doc, "完成度スコア", num, anchor, 1000)
        if ov:
            _para(doc, f"総合 {ov.get('score')}点（{ov.get('stage')}）。採点は証拠の強さで行います。仮定のみは25点、実データの裏付けで50点以上に上がります。", color=MUTE, size=9.5)
        radar = report.get("radar_png")
        if radar:
            pr = doc.add_paragraph()
            pr.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pr.add_run().add_picture(radar, width=Cm(13.5))
        widths = [4.3, 1.8, 10.9]
        t = _table(doc, widths)
        hr = t.add_row()
        _apply_widths(hr, widths)
        for c, txt, al in zip(hr.cells, ["評価軸", "点数", "採点根拠"], [None, WD_ALIGN_PARAGRAPH.CENTER, None]):
            _cell_text(c, txt, bold=True, size=9.5, color=WHITE, align=al)
            _cell_shade(c, NAVY_HEX)
            _cell_margins(c)
        for i, row in enumerate(report["score_table"]):
            rr = t.add_row()
            _apply_widths(rr, widths)
            _cell_text(rr.cells[0], row.get("axis", ""), size=9.5, bold=True, color=NAVY)
            _cell_text(rr.cells[1], str(row.get("score", "")), size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
            _cell_text(rr.cells[2], row.get("basis", ""), size=9)
            for c in rr.cells:
                _cell_margins(c)
                if i % 2 == 1:
                    _cell_shade(c, ZEBRA_HEX)

    def render_risk():
        if not has["risk_table"]:
            return
        num, anchor = _next_bm()
        _section_heading(doc, "リスク登録簿", num, anchor, 1001)
        _para(doc, "致命的リスクが未対応で残る限り、事業計画書には進みません。宿題で状態を動かします。", color=MUTE, size=9.5)
        widths = [1.1, 6.9, 1.9, 1.9, 4.7]
        t = _table(doc, widths)
        hr = t.add_row()
        _apply_widths(hr, widths)
        for c, txt in zip(hr.cells, ["ID", "リスク", "深刻度", "状態", "対応"]):
            _cell_text(c, txt, bold=True, size=9.5, color=WHITE)
            _cell_shade(c, NAVY_HEX)
            _cell_margins(c)
        for i, row in enumerate(report["risk_table"]):
            rr = t.add_row()
            _apply_widths(rr, widths)
            vals = [row.get("id", ""), row.get("risk", ""), row.get("severity", ""), row.get("status", ""), row.get("action", "")]
            for j, (c, v) in enumerate(zip(rr.cells, vals)):
                _cell_text(c, v, size=9, bold=(j == 0), color=NAVY if j == 0 else INK)
                _cell_margins(c)
                if i % 2 == 1:
                    _cell_shade(c, ZEBRA_HEX)

    def render_homework():
        if not has["homework"]:
            return
        doc.add_page_break()
        num, anchor = _next_bm()
        _section_heading(doc, "宿題ドリル（記入式）", num, anchor, 1002)
        _para(doc, "優先度順に並んでいます。上から着手してください。淡い青の欄が、あなたの記入欄です。", color=MUTE, size=9.5)
        default_fields = [
            "実施日", "やったこと・会った相手", "わかった数字・事実",
            "出典・裏付け（資料名、聞いた相手の所属）", "詰まった点・わからなかったこと",
        ]
        for hw in report["homework"]:
            _sub_heading(doc, f"宿題{hw.get('no')}　対象軸: {hw.get('axis', '')}")
            _para(doc, hw.get("title", ""), bold=True, size=10.5, space_after=2)
            if hw.get("how"):
                _para(doc, "やり方: " + hw["how"], size=9.5, color=MUTE, space_after=4)
            widths = [4.5, 12.5]
            t = _table(doc, widths)
            for field in hw.get("fields", default_fields):
                rr = t.add_row()
                _apply_widths(rr, widths)
                _row_height(rr, 1.5)
                _cell_text(rr.cells[0], field, bold=True, size=9.5, color=NAVY)
                _cell_shade(rr.cells[0], BOX_HEX)
                _cell_margins(rr.cells[0])
                _cell_text(rr.cells[1], WRITE_HINT, size=9, color=RGBColor(0xA6, 0xB0, 0xBC))
                _cell_margins(rr.cells[1])

    def render_premise():
        if not has["premise_questions"]:
            return
        num, anchor = _next_bm()
        _section_heading(doc, "前提情報回答シート", num, anchor, 1003)
        _para(doc, "協議で仮定を置いた前提です。分かる範囲で記入してください。空欄でも協議は続きますが、埋まるほど診断の精度が上がります。", color=MUTE, size=9.5)
        widths = [8.5, 8.5]
        t = _table(doc, widths)
        hr = t.add_row()
        _apply_widths(hr, widths)
        for c, txt in zip(hr.cells, ["質問", "回答記入欄"]):
            _cell_text(c, txt, bold=True, size=9.5, color=WHITE)
            _cell_shade(c, NAVY_HEX)
            _cell_margins(c)
        for q in report["premise_questions"]:
            rr = t.add_row()
            _apply_widths(rr, widths)
            _row_height(rr, 1.2)
            _cell_text(rr.cells[0], q, size=9.5)
            _cell_shade(rr.cells[0], BOX_HEX)
            _cell_margins(rr.cells[0])
            _cell_text(rr.cells[1], WRITE_HINT, size=9, color=RGBColor(0xA6, 0xB0, 0xBC))
            _cell_margins(rr.cells[1])

    renderers = {"scorecard": render_scorecard, "risk": render_risk,
                 "homework": render_homework, "premise": render_premise}
    done = set()
    # Bookmark IDs must be unique across the whole document. Marker sections
    # use fixed IDs 1000-1003; generic level-1 sections get their own
    # monotonically increasing counter so two generic sections rendered back
    # to back (i.e. before any marker section bumps len(done)) never collide.
    _generic_bm_id = [1100]

    def _next_generic_bm_id():
        val = _generic_bm_id[0]
        _generic_bm_id[0] += 1
        return val

    # ---- 本文 ----
    for sec in sections:
        marker = sec.get("type")
        if marker in renderers:
            renderers[marker]()
            done.add(marker)
            continue

        level = sec.get("level", 1)
        heading = sec.get("heading", "")
        if heading:
            if level == 1:
                num, anchor = _next_bm()
                _section_heading(doc, heading, num, anchor, _next_generic_bm_id())
            else:
                _sub_heading(doc, heading)

        for para in sec.get("paragraphs", []):
            _para(doc, para)
        for b in sec.get("bullets", []):
            _bullet(doc, b)

        tbl = sec.get("table")
        if tbl and tbl.get("rows"):
            header = tbl.get("header")
            n_cols = len(header or tbl["rows"][0])
            widths = tbl.get("widths") or [17.0 / n_cols] * n_cols
            t = _table(doc, widths)
            if header:
                hr = t.add_row()
                _apply_widths(hr, widths)
                for c, txt in zip(hr.cells, header):
                    _cell_text(c, str(txt), bold=True, size=9.5, color=WHITE)
                    _cell_shade(c, NAVY_HEX)
                    _cell_margins(c)
            for i, row in enumerate(tbl["rows"]):
                rr = t.add_row()
                _apply_widths(rr, widths)
                for j, (c, txt) in enumerate(zip(rr.cells, row)):
                    _cell_text(c, str(txt), size=9, bold=(j == 0), color=NAVY if j == 0 else INK)
                    _cell_margins(c)
                    if i % 2 == 1:
                        _cell_shade(c, ZEBRA_HEX)
            _para(doc, "", space_after=2)

        if sec.get("score_line"):
            st = _table(doc, [17.0])
            rr = st.add_row()
            _apply_widths(rr, [17.0])
            _cell_text(rr.cells[0], sec["score_line"], bold=True, size=9.5, color=NAVY)
            _cell_shade(rr.cells[0], BOX_HEX)
            _cell_margins(rr.cells[0], top=60, bottom=60)

    for key in ("scorecard", "risk", "homework", "premise"):
        if key not in done:
            renderers[key]()

    if not is_plan:
        note = doc.add_paragraph()
        note.paragraph_format.space_before = Pt(14)
        _bottom_border(note, RULE_HEX, size=4)
        _para(
            doc,
            "記入が済んだら、このファイルを保存して、そのまま次回の相談にアップロードしてください。"
            "記入内容をもとに、リスク登録簿の状態更新・完成度スコアの再採点（前回比つきレーダーチャート）・次の宿題をお返しします。",
            bold=True, color=NAVY, space_before=6,
        )

    _page_number_footer(doc)

    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")

    doc.save(output)
    print(f"生成: {output}")


def main(argv=None):
    ap = argparse.ArgumentParser(description="協議録／事業計画書のWordを生成")
    ap.add_argument("--report", required=True)
    ap.add_argument("--output", required=True)
    args = ap.parse_args(argv)
    with open(args.report, encoding="utf-8") as f:
        report = json.load(f)
    build(report, args.output)


if __name__ == "__main__":
    sys.exit(main())
